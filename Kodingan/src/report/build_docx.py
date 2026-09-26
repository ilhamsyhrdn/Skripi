"""Bangun skripsi lengkap sebagai satu berkas .docx dari bab-bab markdown.

Format mengikuti skripsi acuan (skripsi teman pengguna tentang klasifikasi
tumor otak) yang diukur langsung dari PDF-nya:

- A4, margin kiri 4 cm, atas-kanan-bawah 3 cm, Times New Roman 12 pt untuk
  seluruh teks termasuk judul bab dan judul sub-bab.
- Isi rata kiri-kanan, spasi ganda, inden baris pertama 1,27 cm; abstrak
  berspasi tunggal.
- Judul tabel di atas tabel, 11 pt biasa; isi tabel 11 pt; kepala tabel tebal
  berlatar #FFF2CC dan diulang pada setiap halaman; keterangan gambar di bawah
  gambar, 12 pt biasa.
- Nomor halaman di tengah bawah pada halaman awal (romawi) dan pada halaman
  pembuka tiap bab, serta di kanan atas pada halaman isi lainnya.
- Halaman judul tanpa nomor; lembar pengesahan berbingkai garis ganda.

Daftar isi, daftar tabel, dan daftar gambar ditulis sebagai field Word, lalu
diisi dengan nomor halaman sungguhan oleh update_fields_word.ps1.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path("D:/skripsi/Naskah Skripsi")
FIG = Path("D:/skripsi/Kodingan/outputs/figures/final")
OUT = SRC / "Skripsi_Lengkap.docx"

BODY_FONT = "Times New Roman"
MONO_FONT = "Courier New"   # huruf kode pada skripsi acuan
BODY_SIZE = Pt(12)
TABLE_SIZE = Pt(11)
INDENT = Cm(1.27)
TEXT_WIDTH_CM = 14.0          # 21 cm - 4 cm - 3 cm
HEADER_FILL = "FFF2CC"        # warna kepala tabel pada skripsi acuan

JUDUL_ID = [("OPTIMASI ", ""), ("TRANSFER LEARNING", "i"), (" EFFICIENTNET-B0 UNTUK KLASIFIKASI "
            "KANKER PARU-PARU PADA CITRA ", ""), ("COMPUTED TOMOGRAPHY", "i"),
            (" (CT) MENGGUNAKAN ", ""), ("FINE-TUNING", "i"), (", ", ""),
            ("DATA AUGMENTATION", "i"), (", DAN ", ""), ("ENSEMBLE MODEL", "i")]
JUDUL_EN = ("OPTIMIZATION OF EFFICIENTNET-B0 TRANSFER LEARNING FOR LUNG CANCER "
            "CLASSIFICATION ON COMPUTED TOMOGRAPHY (CT) IMAGES USING FINE-TUNING, "
            "DATA AUGMENTATION, AND ENSEMBLE MODEL")

# jalur gambar di markdown -> (berkas gambar final, lebar cm)
FIGURES = {
    "Gambar/Gambar_2.1_Arsitektur_CNN.png": ("gambar_2_1_arsitektur_cnn.png", 15.0),
    "Gambar/Gambar_2.2_Blok_MBConv_EfficientNet.png": ("gambar_2_2_blok_mbconv.png", 15.0),
    "Gambar/Gambar_2.3_Blok_Residual_ResNet50.png": ("gambar_2_3_blok_residual.png", 14.0),
    "Gambar/Gambar_3.1_Diagram_Alur_Penelitian.png": ("gambar_3_1_alur_penelitian.png", 9.5),
    "Gambar/Gambar_3.2_Sampel_Kaggle_Al-Yasriy.png": ("gambar_3_2_sampel_kaggle_alyasriy.png", 9.5),
    "Gambar/Gambar_3.3_Sampel_Kaggle_Rathi.png": ("gambar_3_3_sampel_kaggle_rathi.png", 9.5),
    "Gambar/Gambar_3.4_Sampel_LIDC.png": ("gambar_3_4_sampel_lidc.png", 9.5),
    "Gambar/Gambar_3.5_Use_Case_Diagram.png": ("gambar_3_5_use_case.png", 14.5),
    "Gambar/Gambar_4.1_Kurva_Pelatihan.png": ("gambar_4_1_kurva_pelatihan.png", 14.0),
    "Gambar/Gambar_4.2_Kontribusi_Finetuning.png": ("gambar_4_2_kontribusi_finetuning.png", 14.0),
    "Gambar/Gambar_4.3_Kontribusi_Augmentasi.png": ("gambar_4_3_kontribusi_augmentasi.png", 14.0),
    "Gambar/Gambar_4.4_Kontribusi_Ensemble.png": ("gambar_4_4_kontribusi_ensemble.png", 14.0),
    "Gambar/Gambar_4.5_Confusion_Matrix_Stacking.png": ("gambar_4_5_confusion_matrix.png", 10.0),
    "Gambar/Gambar_4.6_Kurva_ROC_Stacking.png": ("gambar_4_6_kurva_roc.png", 11.0),
    "Gambar/Gambar_4.7_Sweep_Ambang_Keputusan.png": ("gambar_4_7_sweep_ambang.png", 14.0),
    "Gambar/Gambar_4.8_Performa_Per_Sumber.png": ("gambar_4_8_performa_per_sumber.png", 12.5),
    "Gambar/Gambar_4.9_Confusion_Matrix_Validasi_Biner.png": ("gambar_4_cm_validasi_biner.png", 9.0),
    "Gambar/Gambar_4.10_Aplikasi_Beranda.png": ("gambar_4_9_aplikasi_beranda.png", 14.0),
    "Gambar/Gambar_4.11_Aplikasi_Prediksi.png": ("gambar_4_10_aplikasi_prediksi.png", 14.0),
    "Gambar/Gambar_4.12_Aplikasi_Tolak.png": ("gambar_4_11_aplikasi_tolak.png", 14.0),
}


# ------------------------------------------------------------------ utilitas
def _font(run, size=BODY_SIZE, bold=None, italic=None, name=BODY_FONT):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def _field(run, instr_text):
    """Sisipkan field sederhana (PAGE, TOC, dan sejenisnya) ke dalam run."""
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = instr_text
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = " "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, txt, end):
        run._r.append(el)


# ------------------------------------------------------------------- dokumen
def setup_document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 2.0

    # Gaya judul bawaan dipakai agar field daftar isi dapat membacanya.
    for name, align, before, after in (
        ("Heading 1", WD_ALIGN_PARAGRAPH.CENTER, 0, 24),
        ("Heading 2", WD_ALIGN_PARAGRAPH.LEFT, 6, 0),
        ("Heading 3", WD_ALIGN_PARAGRAPH.LEFT, 6, 0),
    ):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = BODY_SIZE
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        st.element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        st.element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        st.paragraph_format.alignment = align
        st.paragraph_format.line_spacing = 2.0
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.keep_with_next = True

    # Gaya entri daftar isi: Times New Roman, spasi 1,5, tanpa jarak antar entri.
    for lvl in (1, 2, 3):
        try:
            st = doc.styles[f"TOC {lvl}"]
        except KeyError:
            st = doc.styles.add_style(f"TOC {lvl}", 1)
            st.base_style = normal
        st.font.name = BODY_FONT
        st.font.size = BODY_SIZE
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.left_indent = Cm({1: 0, 2: 0.6, 3: 1.4}[lvl])

    # Keterangan gambar 12 pt dan judul tabel 11 pt, masing-masing gaya
    # tersendiri agar daftar gambar dan daftar tabel terbentuk terpisah.
    for name, size in (("Caption Gambar", BODY_SIZE), ("Caption Tabel", TABLE_SIZE)):
        st = doc.styles.add_style(name, 1)
        st.base_style = normal
        st.font.name = BODY_FONT
        st.font.size = size
        st.font.bold = False
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.first_line_indent = Cm(0)
        st.quick_style = False

    _page_setup(doc.sections[0])
    return doc


def _page_setup(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(4.0)
    section.right_margin = Cm(3.0)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def _clear(part):
    for p in part.paragraphs:
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
    return part.paragraphs[0]


def _page_field_paragraph(part, align):
    p = _clear(part)
    p.alignment = align
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Cm(0)
    _field(_font(p.add_run()), " PAGE ")


_SESUDAH_PGNUM = ("cols", "formProt", "vAlign", "noEndnote", "titlePg", "textDirection",
                  "bidi", "rtlGutter", "docGrid", "printerSettings")


def _sisip_sectpr(sectPr, el, sesudahnya):
    """Sisipkan el tepat sebelum anak sectPr pertama yang menurut skema OOXML
    harus berada setelahnya."""
    for child in sectPr:
        if child.tag.split("}")[-1] in sesudahnya:
            child.addprevious(el)
            return
    sectPr.append(el)


def _pg_num(section, fmt, start=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        _sisip_sectpr(sectPr, pg, _SESUDAH_PGNUM)
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    elif pg.get(qn("w:start")) is not None:
        del pg.attrib[qn("w:start")]


def _unlink(section):
    for part in (section.header, section.footer, section.first_page_header,
                 section.first_page_footer):
        part.is_linked_to_previous = False


def number_front(section, start=None, show=True):
    """Halaman awal: nomor romawi di tengah bawah (atau disembunyikan)."""
    _unlink(section)
    section.different_first_page_header_footer = False
    _clear(section.header)
    if show:
        _page_field_paragraph(section.footer, WD_ALIGN_PARAGRAPH.CENTER)
    else:
        _clear(section.footer)
    _pg_num(section, "lowerRoman", start)


def number_chapter(section, start=None):
    """Halaman bab: tengah bawah di halaman pembuka, kanan atas di halaman lain."""
    _unlink(section)
    section.different_first_page_header_footer = True
    _clear(section.first_page_header)
    _page_field_paragraph(section.first_page_footer, WD_ALIGN_PARAGRAPH.CENTER)
    _page_field_paragraph(section.header, WD_ALIGN_PARAGRAPH.RIGHT)
    _clear(section.footer)
    _pg_num(section, "decimal", start)


def new_section(doc):
    """Section baru mewarisi seluruh pengaturan section sebelumnya, termasuk
    bingkai halaman, sehingga bingkai lembar pengesahan harus dibuang di sini
    agar tidak ikut muncul pada halaman-halaman berikutnya."""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    _page_setup(s)
    for b in s._sectPr.findall(qn("w:pgBorders")):
        s._sectPr.remove(b)
    return s


def page_border(section):
    """Bingkai garis ganda (tebal di luar, tipis di dalam) seperti lembar
    pengesahan skripsi acuan."""
    sectPr = section._sectPr
    b = OxmlElement("w:pgBorders")
    b.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "thickThinSmallGap")
        el.set(qn("w:sz"), "24")
        el.set(qn("w:space"), "24")
        el.set(qn("w:color"), "000000")
        b.append(el)
    _sisip_sectpr(sectPr, b, ("lnNumType", "pgNumType") + _SESUDAH_PGNUM)


# ------------------------------------------------------------ teks sebaris
def _tokens(text):
    """Pecah teks menjadi potongan (isi, tebal, miring, kode).

    Ditulis sebagai pemindai sederhana, bukan satu regex, supaya teks tebal
    yang memuat kata miring seperti **ResNet50 *fold* 4** terbaca benar.
    """
    out, buf = [], []
    bold = italic = False
    i, n = 0, len(text)

    def flush():
        if buf:
            out.append(("".join(buf), bold, italic, False))
            buf.clear()

    while i < n:
        c = text[i]
        if c == "`":
            j = text.find("`", i + 1)
            if j > i:
                flush()
                out.append((text[i + 1:j], bold, italic, True))
                i = j + 1
                continue
        if c == "$" and i + 1 < n and text[i + 1] != " ":
            j = text.find("$", i + 1)
            if j > i + 1:
                flush()
                out.append((text[i + 1:j], bold, True, False))
                i = j + 1
                continue
        if text.startswith("**", i):
            flush(); bold = not bold; i += 2
            continue
        if c == "*" and (i + 1 < n and text[i + 1] != " " or italic):
            flush(); italic = not italic; i += 1
            continue
        buf.append(c)
        i += 1
    flush()
    return out


def add_runs(paragraph, text, size=BODY_SIZE, force_bold=None):
    text = text.replace(" -- ", " – ")
    for isi, b, it, kode in _tokens(text):
        if not isi:
            continue
        r = paragraph.add_run(isi)
        if kode:
            _font(r, size=Pt(size.pt - 1.5), name=MONO_FONT)
        else:
            _font(r, size=size, bold=b if force_bold is None else (force_bold or b),
                  italic=it)
    return paragraph


# ----------------------------------------------------------------- paragraf
def body_paragraph(doc, text, indent=True, spacing=2.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT if indent else Cm(0)
    add_runs(p, text)
    return p


def caption(doc, text, style):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    if style == "Caption Tabel":
        pf.space_before = Pt(12); pf.space_after = Pt(4); pf.keep_with_next = True
        size = TABLE_SIZE
    else:
        pf.space_before = Pt(4); pf.space_after = Pt(12)
        size = BODY_SIZE
    add_runs(p, text, size=size)
    return p


def chapter_heading(doc, bab, title):
    """Satu paragraf Heading 1 berisi dua baris ("BAB I" lalu judulnya)."""
    p = doc.add_paragraph(style="Heading 1")
    _font(p.add_run(bab.upper()), bold=True)
    if title:
        p.runs[-1].add_break()
        add_runs(p, title.upper(), force_bold=True)
    return p


def front_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    add_runs(p, text.upper(), force_bold=True)
    return p


def section_heading(doc, text, level=2):
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    pf = p.paragraph_format
    pf.left_indent = Cm(0)
    pf.first_line_indent = Cm(0)
    m = re.match(r"^((?:\d+\.)+\d*)\s+(.*)$", text)
    if m:
        num, rest = m.groups()
        pf.tab_stops.add_tab_stop(Cm(1.0 if level == 2 else 1.4))
        _font(p.add_run(f"{num}\t"), bold=True)
        add_runs(p, rest, force_bold=True)
    else:
        add_runs(p, text, force_bold=True)
    return p


def add_list_item(doc, marker, text, level=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    pf.left_indent = Cm(0.8 + level * 0.8)
    pf.first_line_indent = Cm(-0.8)
    pf.tab_stops.add_tab_stop(Cm(0.8 + level * 0.8))
    add_runs(p, f"{marker}\t{text}")
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.left_indent = Cm(0.3)
    pf.right_indent = Cm(0.2)
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(12)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "6"); el.set(qn("w:color"), "808080")
        borders.append(el)
    pPr.append(borders)
    for k, line in enumerate(lines):
        r = _font(p.add_run(line), size=Pt(10), name=MONO_FONT)
        if k < len(lines) - 1:
            r.add_break()
    return p


# ------------------------------------------------------------- gambar/rumus
_MATH_DIR = Path("D:/skripsi/Kodingan/outputs/figures/math")


def render_math(latex: str, index: int):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    _MATH_DIR.mkdir(parents=True, exist_ok=True)
    out = _MATH_DIR / f"eq_{index:02d}.png"
    # mathtext paham \frac, \sum, \times; \text{-} dan \qquad perlu dibantu
    body = latex.strip().replace(r"\text{-}", "-").replace(r"\qquad", r"\ \ \ \ \ ")
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${body}$", fontsize=16)
    fig.savefig(out, dpi=220, bbox_inches="tight", pad_inches=0.06, transparent=False,
                facecolor="white")
    plt.close(fig)
    return out


def add_equation(doc, latex, index):
    path = render_math(latex, index)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.0; pf.space_before = Pt(6); pf.space_after = Pt(6)
    pf.first_line_indent = Cm(0)
    from PIL import Image
    with Image.open(path) as im:
        w_cm = min(10.0, max(3.2, im.width / 220 * 2.54))
    p.add_run().add_picture(str(path), width=Cm(w_cm))


def add_figure(doc, md_path, cap_text):
    fname, width_cm = FIGURES[md_path]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.0; pf.space_before = Pt(12); pf.space_after = Pt(0)
    pf.first_line_indent = Cm(0); pf.keep_with_next = True
    p.add_run().add_picture(str(FIG / fname), width=Cm(width_cm))
    caption(doc, cap_text, "Caption Gambar")


# ------------------------------------------------------------------- tabel
CELL_IMG_RE = re.compile(r"^!\[([\d.,]*)\]\((.+?)\)$")


def _set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _row_props(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit"); trPr.append(cant)
    if header:
        th = OxmlElement("w:tblHeader"); trPr.append(th)


def _plain(cell):
    return re.sub(r"[*`]", "", cell).strip()


_FONT_FILES = {(False, False): "times.ttf", (True, False): "timesbd.ttf",
               (False, True): "timesi.ttf", (True, True): "timesbi.ttf"}
_FONT_CACHE: dict = {}


def _lebar_cm(teks, size_pt, bold=False, italic=False):
    """Lebar teks sebenarnya dalam cm, diukur dari berkas huruf Times New Roman."""
    from PIL import ImageFont
    key = (bold, italic)
    if key not in _FONT_CACHE:
        _FONT_CACHE[key] = ImageFont.truetype(f"C:/Windows/Fonts/{_FONT_FILES[key]}", 100)
    return _FONT_CACHE[key].getlength(teks) / 100 * size_pt * 0.03528


def _lebar_sel(teks, size_pt, header=False, utuh=False):
    """Lebar minimum sel: kata terpanjang (bila boleh dibungkus) atau seluruh
    isi (bila sel pendek yang tidak boleh terpotong, misalnya angka)."""
    total = 0.0
    kata_max = 0.0
    for isi, b, it, kode in _tokens(teks):
        b = b or header
        for w in re.split(r"(\s+)", isi):
            if not w:
                continue
            lw = _lebar_cm(w, size_pt, b, it)
            total += lw
            if not w.isspace():
                kata_max = max(kata_max, lw)
    return total if utuh else kata_max


PAD_CM = 0.5   # margin dalam sel kiri-kanan ditambah ruang aman


def _column_widths(rows, size_pt):
    """Kolom angka dan kolom berisi teks pendek diberi lebar secukupnya agar
    isinya tidak terpotong; sisa lebar halaman dibagikan kepada kolom berisi
    teks panjang sebanding panjang teksnya."""
    ncol = len(rows[0])
    need, bobot = [], []
    for j in range(ncol):
        col = [r[j] if j < len(r) else "" for r in rows]
        imgs = [CELL_IMG_RE.match(c.strip()) for c in col[1:]]
        if any(imgs):
            w = max(float(m.group(1).replace(",", ".") or 3.2) for m in imgs if m)
            need.append(max(w + PAD_CM, _lebar_sel(col[0], size_pt, header=True) + PAD_CM))
            bobot.append(0.0)
            continue
        pendek = _is_center_col(rows, j)
        # angka tidak boleh terpotong, sedangkan teks pendek seperti
        # "100% Berhasil" boleh terbungkus dua baris seperti pada skripsi acuan
        angka = _is_numeric_col(rows, j)
        n_hdr = _lebar_sel(col[0], size_pt, header=True)
        n_body = max((_lebar_sel(c, size_pt, utuh=angka) for c in col[1:]), default=0)
        need.append(max(n_hdr, n_body) + PAD_CM)
        bobot.append(0.0 if pendek else
                     sum(_lebar_sel(c, size_pt, utuh=True) for c in col[1:]) / max(len(col) - 1, 1))
    lebar = list(need)
    sisa = TEXT_WIDTH_CM - sum(lebar)
    if sisa > 0:
        tot = sum(bobot)
        if tot > 0:
            lebar = [w + sisa * b / tot for w, b in zip(lebar, bobot)]
        else:
            lebar = [w + sisa / ncol for w in lebar]
    return lebar, sum(need) <= TEXT_WIDTH_CM + 1e-6


def _is_numeric_col(rows, j):
    body = [_plain(r[j]) for r in rows[1:] if j < len(r)]
    return bool(body) and all(re.fullmatch(r"[\d.,%\-+−×\s/()e]*", c) and len(c) <= 18 for c in body)


def _is_center_col(rows, j):
    body = [_plain(r[j]) for r in rows[1:] if j < len(r)]
    if not body:
        return False
    numeric = all(re.fullmatch(r"[\d.,%\-+−×\s/()e]*", c) and len(c) <= 18 for c in body)
    short = all(len(c) <= 16 for c in body)
    return numeric or short


def add_table(doc, rows):
    header, *body = rows
    ncol = len(header)
    size = TABLE_SIZE
    widths, muat = _column_widths(rows, size.pt)
    if not muat:
        # tabel dengan banyak kolom: huruf diperkecil satu tingkat agar tidak
        # ada kata yang terpotong di tengah
        size = Pt(10)
        widths, muat = _column_widths(rows, size.pt)
        print(f"  [tabel 10 pt] {' | '.join(_plain(h)[:14] for h in header)}", flush=True)
        if not muat:
            skala = TEXT_WIDTH_CM / sum(widths)
            widths = [w * skala for w in widths]
            print(f"  [PERINGATAN: tabel masih terlalu lebar, diskala {skala:.2f}]", flush=True)
    center = [_is_center_col(rows, j) for j in range(ncol)]
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    for j, w in enumerate(widths):
        for cell in table.columns[j].cells:
            cell.width = Cm(w)

    for i, row in enumerate(rows):
        _row_props(table.rows[i], header=(i == 0))
        for j in range(ncol):
            text = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = ""
            cell.vertical_alignment = (WD_CELL_VERTICAL_ALIGNMENT.CENTER if i == 0
                                       else WD_CELL_VERTICAL_ALIGNMENT.TOP)
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.line_spacing = 1.0; pf.space_after = Pt(1); pf.space_before = Pt(1)
            pf.first_line_indent = Cm(0)
            m = CELL_IMG_RE.match(text.strip())
            if m and i > 0:
                w_cm = float(m.group(1).replace(",", ".") or 3.2)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(SRC / m.group(2)), width=Cm(w_cm))
            else:
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_runs(p, text, size=size, force_bold=True)
                else:
                    # rata kiri-kanan hanya untuk paragraf panjang seperti kolom
                    # analisis; sel pendek rata kiri agar tidak muncul celah lebar
                    paragraf = len(_plain(text)) > 150
                    p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if center[j]
                                   else WD_ALIGN_PARAGRAPH.JUSTIFY if paragraf
                                   else WD_ALIGN_PARAGRAPH.LEFT)
                    add_runs(p, text, size=size)
            _set_cell_border(cell)
            if i == 0:
                _shade(cell, HEADER_FILL)
    gap = doc.add_paragraph()
    gap.paragraph_format.line_spacing = 1.0
    gap.paragraph_format.first_line_indent = Cm(0)
    return table


# ---------------------------------------------------------- parser markdown
CAPTION_TAB_RE = re.compile(r"^(Tabel\s+[\d.]+\s+.*)$")
IMG_RE = re.compile(r"^!\[(.*?)\]\((.*?)\)$")
OL_RE = re.compile(r"^(\d+)\.\s+(.*)$")
AL_RE = re.compile(r"^([a-z])\.\s+(.*)$")
UL_RE = re.compile(r"^[-*]\s+(.*)$")
_EQ = {"n": 0}


def parse_markdown(doc, path: Path, on_chapter, pustaka=False):
    """on_chapter(bab, judul) dipanggil sebelum judul bab ditulis, agar setiap
    bab dapat dibuka pada section baru dengan penomoran halamannya sendiri."""
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    pending_bab = None
    pending_table_caption = None
    pending_figure = None

    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue

        if line.startswith("# "):
            text = line[2:].strip()
            if re.match(r"^BAB\s+[IVX]+$", text):
                pending_bab = text
                i += 1
                continue
            on_chapter()
            if pending_bab:
                chapter_heading(doc, pending_bab, text)
                pending_bab = None
            else:
                front_heading(doc, text)
            i += 1
            continue

        if line.startswith("### "):
            section_heading(doc, line[4:].strip(), level=3); i += 1; continue
        if line.startswith("## "):
            section_heading(doc, line[3:].strip(), level=2); i += 1; continue

        m = IMG_RE.match(s)
        if m:
            if m.group(2) not in FIGURES:
                raise KeyError(f"{path.name}: gambar belum terdaftar di FIGURES -> {m.group(2)}")
            pending_figure = m.group(2); i += 1; continue

        if pending_figure:
            cap = s.strip("*")
            if cap.startswith("Gambar"):
                add_figure(doc, pending_figure, cap)
                pending_figure = None
                i += 1
                continue
            pending_figure = None

        if CAPTION_TAB_RE.match(s):
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and lines[j].lstrip().startswith("|"):
                pending_table_caption = s
                i += 1
                continue

        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if pending_table_caption:
                caption(doc, pending_table_caption, "Caption Tabel")
                pending_table_caption = None
            if rows:
                add_table(doc, rows)
            continue

        if s.startswith("$$"):
            if s.endswith("$$") and len(s) > 4:
                latex = s[2:-2]; i += 1
            else:
                parts = [s[2:]]; i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    parts.append(lines[i]); i += 1
                if i < len(lines):
                    parts.append(lines[i].strip()[:-2]); i += 1
                latex = " ".join(parts)
            _EQ["n"] += 1
            add_equation(doc, latex, _EQ["n"])
            continue

        if s.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            add_code_block(doc, code)
            continue

        if set(s) <= {"-"} and len(s) >= 3:
            i += 1
            continue

        m = OL_RE.match(s)
        if m:
            add_list_item(doc, f"{m.group(1)}.", m.group(2)); i += 1; continue
        m = AL_RE.match(s)
        if m:
            add_list_item(doc, f"{m.group(1)}.", m.group(2)); i += 1; continue
        m = UL_RE.match(s)
        if m and not s.startswith("**"):
            add_list_item(doc, "\u2022", m.group(1)); i += 1; continue

        # label cuplikan kode ditulis rata kiri tanpa inden, seperti acuan
        if re.match(r"^\*script\* kode program", s):
            body_paragraph(doc, s, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
            i += 1
            continue

        if pustaka:
            # daftar pustaka acuan: inden gantung 0,85 cm, spasi tunggal,
            # jarak 6 pt antar-entri
            pr = body_paragraph(doc, s, indent=False, spacing=1.0)
            pr.paragraph_format.left_indent = Cm(0.85)
            pr.paragraph_format.first_line_indent = Cm(-0.85)
            pr.paragraph_format.space_after = Pt(6)
        else:
            body_paragraph(doc, s)
        i += 1


# ----------------------------------------------------------- halaman awal
def _centred(doc, text_or_runs, bold=False, italic=False, spacing=1.5, after=0, before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = spacing; pf.space_after = Pt(after); pf.space_before = Pt(before)
    pf.first_line_indent = Cm(0)
    runs = text_or_runs if isinstance(text_or_runs, list) else [(text_or_runs, "")]
    for teks, gaya in runs:
        _font(p.add_run(teks), bold=bold, italic=italic or "i" in gaya)
    return p


def title_page(doc):
    for _ in range(1):
        _centred(doc, "", spacing=1.0)
    _centred(doc, JUDUL_ID, bold=True, after=48)
    _centred(doc, "SKRIPSI", bold=True, after=48)
    _centred(doc, "diajukan untuk menempuh ujian sarjana")
    _centred(doc, "pada Fakultas [nama fakultas]")
    _centred(doc, "Universitas [nama universitas]", after=48)
    _centred(doc, "[NAMA LENGKAP MAHASISWA]")
    _centred(doc, "NPM [NOMOR POKOK MAHASISWA]", after=60)
    _centred(doc, "[LOGO UNIVERSITAS]", italic=True, after=90)
    for t in ("UNIVERSITAS [NAMA UNIVERSITAS]", "FAKULTAS [NAMA FAKULTAS]",
              "PROGRAM STUDI [NAMA PROGRAM STUDI]", "[KOTA]", "2026"):
        _centred(doc, t, spacing=1.0)


def lembar_pengesahan(doc):
    _centred(doc, "", spacing=1.0)
    _centred(doc, "SKRIPSI", bold=True, after=12)
    _centred(doc, JUDUL_ID, bold=True, after=18)
    _centred(doc, JUDUL_EN, bold=True, italic=True, after=18)
    _centred(doc, "Telah dipersiapkan dan disusun oleh", after=12)
    _centred(doc, "[NAMA LENGKAP MAHASISWA]", spacing=1.2)
    _centred(doc, "NPM [NOMOR POKOK MAHASISWA]", after=12)
    _centred(doc, "Telah dipertahankan di depan Tim Penguji", spacing=1.2)
    _centred(doc, "pada tanggal [tanggal sidang]", after=12)
    _centred(doc, "Susunan Tim Penguji", after=12)

    anggota = [("[Nama Ketua Tim Penguji]", "Ketua Tim Penguji"),
               ("[Nama Pembimbing Utama]", "Pembimbing"),
               ("[Nama Pembimbing Pendamping]", "Co-Pembimbing"),
               ("[Nama Penguji]", "Penguji"),
               ("[Nama Penguji]", "Penguji")]
    t = doc.add_table(rows=len(anggota), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, w in enumerate((0.9, 6.0, 3.8, 3.3)):
        for cell in t.columns[j].cells:
            cell.width = Cm(w)
    for i, (nama, peran) in enumerate(anggota):
        isi = (f"{i + 1}.", None, peran, "..........................")
        for j in range(4):
            p = t.cell(i, j).paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.first_line_indent = Cm(0)
            if j == 1:
                _font(p.add_run(nama)).underline = True
                _font(p.add_run()).add_break()
                _font(p.add_run("NIP. [NIP]"))
            elif j == 3:
                _font(p.add_run()).add_break()
                _font(p.add_run(isi[j]))
            else:
                _font(p.add_run(isi[j]))


def kata_pengantar(doc):
    front_heading(doc, "Kata Pengantar")
    p = body_paragraph(doc, "Puji dan syukur penulis panjatkan ke hadirat Tuhan Yang Maha Esa, "
                            "yang telah melimpahkan rahmat dan karunia-Nya sehingga penulis dapat "
                            "menyelesaikan penyusunan skripsi yang berjudul “")
    for teks, gaya in JUDUL_ID:
        _font(p.add_run(teks), bold=True, italic="i" in gaya)
    add_runs(p, "” sebagai salah satu syarat menempuh ujian sarjana pada Program Studi "
                "[nama program studi], Fakultas [nama fakultas], Universitas [nama universitas].")
    body_paragraph(doc, "Dalam proses penyusunan dan penulisan skripsi ini tidak terlepas dari "
                        "bantuan, bimbingan, serta dukungan dari berbagai pihak. Oleh karena itu, "
                        "dalam kesempatan ini penulis mengucapkan terima kasih sebanyak-banyaknya "
                        "kepada:")
    items = [
        "[Nama Dekan], selaku Dekan Fakultas [nama fakultas] Universitas [nama universitas].",
        "[Nama Kepala Departemen], selaku Kepala Departemen [nama departemen] Fakultas [nama "
        "fakultas] Universitas [nama universitas].",
        "[Nama Ketua Program Studi], selaku Ketua Program Studi [nama program studi] Fakultas "
        "[nama fakultas] Universitas [nama universitas].",
        "[Nama Pembimbing Utama], selaku dosen pembimbing utama yang telah meluangkan waktu, "
        "arahan, dan koreksi selama proses penyusunan skripsi ini.",
        "[Nama Pembimbing Pendamping], selaku dosen pembimbing pendamping.",
        "Seluruh dosen dan staf Program Studi [nama program studi] yang telah memberikan ilmu "
        "dan bantuan administratif selama masa perkuliahan.",
        "Keluarga penulis yang selalu memberikan motivasi dan doa yang menjadi pendorong dalam "
        "penyelesaian skripsi ini.",
        "Rekan-rekan mahasiswa yang telah memberikan bantuan dan diskusi selama proses "
        "penelitian berlangsung.",
    ]
    for k, t in enumerate(items, start=1):
        add_list_item(doc, f"{k}.", t)
    body_paragraph(doc, "Penulis menyadari bahwa skripsi ini masih memiliki kekurangan, sehingga "
                        "kritik dan saran yang membangun sangat penulis harapkan. Semoga skripsi "
                        "ini dapat bermanfaat bagi pembaca dan bagi pengembangan penelitian "
                        "selanjutnya.")
    for t in ("[Kota], [tanggal] 2026", "", "", "Penulis"):
        body_paragraph(doc, t, indent=False, align=WD_ALIGN_PARAGRAPH.RIGHT)


def abstrak(doc, front, key, head):
    front_heading(doc, head)
    block = front.split(key, 1)[1].split("---", 1)[0]
    for para in [x.strip() for x in block.strip().splitlines() if x.strip()]:
        is_kw = para.startswith("**Kata Kunci**") or para.startswith("**Keywords**")
        p = body_paragraph(doc, para, indent=not is_kw, spacing=1.0)
        if is_kw:
            p.paragraph_format.space_before = Pt(12)
        else:
            p.paragraph_format.space_after = Pt(0)


def toc_list(doc, head, instr):
    front_heading(doc, head)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(0)
    _field(_font(p.add_run()), instr)


def daftar_lampiran(doc):
    front_heading(doc, "Daftar Lampiran")
    body_paragraph(doc, "Lampiran 1 Tautan Repositori GitHub Berisi Kode Sumber",
                   indent=False, spacing=1.5, align=WD_ALIGN_PARAGRAPH.LEFT)


def riwayat_hidup(doc):
    front_heading(doc, "Riwayat Hidup")
    body_paragraph(doc, "[Bagian ini diisi dengan riwayat hidup penulis: nama lengkap, tempat "
                        "dan tanggal lahir, riwayat pendidikan formal dari jenjang dasar hingga "
                        "perguruan tinggi, serta pengalaman organisasi, kegiatan, atau prestasi "
                        "yang relevan.]")


# ----------------------------------------------------------------- rakitan
def main():
    doc = setup_document()

    # halaman judul: dihitung sebagai halaman i tetapi nomornya tidak tampil
    title_page(doc)
    number_front(doc.sections[0], start=1, show=False)

    # lembar pengesahan berbingkai, halaman ii
    s = new_section(doc)
    number_front(s)
    page_border(s)
    lembar_pengesahan(doc)

    front = (SRC / "00_Halaman_Depan_dan_Abstrak.md").read_text(encoding="utf-8")
    s = new_section(doc)
    number_front(s)
    kata_pengantar(doc)
    for key, head in (("## ABSTRAK", "Abstrak"), ("## ABSTRACT", "*Abstract*")):
        doc.add_page_break()
        abstrak(doc, front, key, head)
    doc.add_page_break(); toc_list(doc, "Daftar Isi", r' TOC \o "1-3" \h \z \u ')
    doc.add_page_break(); toc_list(doc, "Daftar Tabel", r' TOC \h \z \t "Caption Tabel,1" ')
    doc.add_page_break(); toc_list(doc, "Daftar Gambar", r' TOC \h \z \t "Caption Gambar,1" ')
    doc.add_page_break(); daftar_lampiran(doc)

    state = {"first": True}

    def on_chapter():
        s = new_section(doc)
        number_chapter(s, start=1 if state["first"] else None)
        state["first"] = False

    for name in ("BAB_I_Pendahuluan.md", "BAB_II_Tinjauan_Pustaka.md",
                 "BAB_III_Analisis_dan_Perancangan.md", "BAB_IV_Hasil_dan_Pembahasan.md",
                 "BAB_V_Kesimpulan_dan_Saran.md", "DAFTAR_PUSTAKA.md", "LAMPIRAN.md"):
        parse_markdown(doc, SRC / name, on_chapter, pustaka=(name == "DAFTAR_PUSTAKA.md"))

    on_chapter()
    riwayat_hidup(doc)

    try:
        doc.save(OUT)
        print(f"Saved: {OUT}", flush=True)
    except PermissionError:
        alt = OUT.with_name(OUT.stem + "_BARU.docx")
        doc.save(alt)
        print(f"Saved: {alt}  (berkas utama terkunci Word)", flush=True)


if __name__ == "__main__":
    main()
